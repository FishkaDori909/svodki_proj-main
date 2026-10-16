"""Проверки безопасности и приёма почты. Запуск: python -m unittest discover tests"""
import mailbox
import os
import tempfile
import textwrap
import unittest
import zipfile
from email.message import EmailMessage
from pathlib import Path

from openpyxl import Workbook

from svodka import mail_in, mail_out, watcher
from svodka.builder import RunState, put_row, set_cell
from svodka.config import load_config
from svodka.errors import ConfigError, SvodkaError
from svodka.mail_in import sender_allowed
from svodka.safeio import check_source_file, safe_filename

CONFIG = textwrap.dedent("""
    [paths]
    inbox = "{root}/inbox"
    results = "{root}/results"
    errors = "{root}/errors"
    state = "{root}/state"
    log_file = "{root}/log.txt"
    template_file = "templates/svodka_template.xlsx"
    [mail_in]
    enabled = true
    mbox_path = "{root}/box"
    allowed_senders = ["good@company.ru", "@trusted.ru"]
""")


def make_cfg(root: Path, extra=""):
    path = root / "config.toml"
    path.write_text(CONFIG.format(root=root.as_posix()) + extra, encoding="utf-8")
    return load_config(path)


def xlsx_bytes(root: Path) -> bytes:
    p = root / "tmp.xlsx"
    Workbook().save(p)
    return p.read_bytes()


class SafeFilename(unittest.TestCase):
    def test_path_traversal_is_stripped(self):
        self.assertEqual(safe_filename("../../etc/evil.xlsx"), "evil.xlsx")
        self.assertEqual(safe_filename("..\\..\\win\\evil.xlsb"), "evil.xlsb")

    def test_bad_extension_rejected(self):
        self.assertIsNone(safe_filename("run.exe"))
        self.assertIsNone(safe_filename("report.xlsx.exe"))
        self.assertIsNone(safe_filename(""))
        self.assertIsNone(safe_filename(".xlsx"))

    def test_control_chars_removed(self):
        self.assertEqual(safe_filename("a\r\nb.xlsx"), "ab.xlsx")

    def test_normal_name_kept(self):
        self.assertEqual(safe_filename("Сводка ПОНМ октябрь.xlsb"), "Сводка ПОНМ октябрь.xlsb")


class Senders(unittest.TestCase):
    allowed = ("good@company.ru", "@trusted.ru")

    def test_rules(self):
        self.assertTrue(sender_allowed("Good@Company.ru", self.allowed))
        self.assertTrue(sender_allowed("anyone@trusted.ru", self.allowed))
        self.assertFalse(sender_allowed("evil@company.ru", self.allowed))
        self.assertFalse(sender_allowed("x@evil-trusted.ru", self.allowed))
        self.assertFalse(sender_allowed("", self.allowed))


class FormulaInjection(unittest.TestCase):
    def test_text_from_source_never_becomes_formula(self):
        ws = Workbook().active
        run = RunState()
        put_row(ws, 1, ['=HYPERLINK("http://evil","x")', "+1+1", "ok"], run)
        self.assertEqual(ws["D1"].data_type, "s")
        self.assertEqual(ws["D1"].value, '=HYPERLINK("http://evil","x")')
        set_cell(ws, "A1", "=1+1", run)               # формулы самой программы остаются формулами
        self.assertEqual(ws["A1"].data_type, "f")


class SourceChecks(unittest.TestCase):
    def test_not_a_zip_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "x.xlsx"
            p.write_text("это не excel")
            with self.assertRaises(SvodkaError):
                check_source_file(p, 50, 500)

    def test_big_unpacked_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "bomb.xlsx"
            with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
                z.writestr("a.xml", b"0" * (3 * 1024 * 1024))
            with self.assertRaises(SvodkaError):
                check_source_file(p, 50, 2)


class Config(unittest.TestCase):
    def test_missing_file(self):
        with self.assertRaises(ConfigError):
            load_config(Path("нет-такого.toml"))

    def test_typo_in_key_is_reported(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "c.toml"
            p.write_text("[watch]\ninterval_secondz = 5\n", encoding="utf-8")
            with self.assertRaises(ConfigError):
                load_config(p)

    def test_mail_in_requires_allowlist(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "c.toml"
            p.write_text('[mail_in]\nenabled = true\nmbox_path = "x"\n', encoding="utf-8")
            with self.assertRaises(ConfigError):
                load_config(p)

    def test_plain_smtp_with_password_over_network_refused(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "c.toml"
            p.write_text('[mail_out]\nenabled = true\nfrom = "a@b.ru"\nto = ["c@d.ru"]\n'
                         '[smtp]\nserver = "mail.example.com"\nsecurity = "none"\n'
                         'login = "u"\npassword = "p"\n', encoding="utf-8")
            with self.assertRaises(ConfigError):
                load_config(p)

    def test_password_from_environment(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "c.toml"
            p.write_text('[mail_out]\nenabled = true\nfrom = "a@b.ru"\nto = ["c@d.ru"]\n'
                         '[smtp]\nserver = "mail.example.com"\nlogin = "u"\n', encoding="utf-8")
            os.environ["SVODKA_SMTP_PASSWORD"] = "secret"
            try:
                self.assertEqual(load_config(p).mail_out.smtp.password, "secret")
            finally:
                del os.environ["SVODKA_SMTP_PASSWORD"]


def add_mail(box, sender, subject, attach_name=None, data=b"", extra_headers=None):
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"] = sender, "svodka@company.ru", subject
    msg["Message-ID"] = f"<{abs(hash((sender, subject, attach_name)))}@test>"
    for k, v in (extra_headers or {}).items():
        msg[k] = v
    msg.set_content("текст")
    if attach_name:
        msg.add_attachment(data, maintype="application", subtype="octet-stream", filename=attach_name)
    box.add(msg)


class Mailbox(unittest.TestCase):
    def test_only_allowed_senders_and_safe_names(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            cfg = make_cfg(root)
            data = xlsx_bytes(root)
            box = mailbox.mbox(str(root / "box"))
            add_mail(box, "Good <good@company.ru>", "ПОНМ", "../../evil.xlsx", data)
            add_mail(box, "Bad <evil@evil.ru>", "ПОНМ", "bad.xlsx", data)
            add_mail(box, "Good <good@company.ru>", "ПОНМ", "program.exe", b"MZ")
            add_mail(box, "Good <good@company.ru>", "Сводка 2-х час ГПН-Ямал — ПОНМ", "loop.xlsx", data)
            add_mail(box, "Good <good@company.ru>", "ПОНМ", "loop2.xlsx", data, {"X-Svodka-Auto": "1"})
            add_mail(box, "Anyone <a@trusted.ru>", "ПОНМ", "domain.xlsx", data)
            box.flush()
            box.close()

            mail_in.scan_mailbox(cfg)

            saved = sorted(p.name for p in cfg.paths.inbox.iterdir())
            self.assertEqual(saved, ["domain.xlsx", "evil.xlsx"])           # evil.xlsx — уже без ../../
            self.assertFalse((root.parent / "evil.xlsx").exists())
            # адрес отправителя запомнен в папке состояния, а не рядом с файлом
            self.assertTrue((cfg.paths.state / "evil.xlsx.json").exists())

            before = sorted(p.name for p in cfg.paths.inbox.iterdir())
            mail_in.scan_mailbox(cfg)                                        # повторный проход ничего не добавляет
            self.assertEqual(before, sorted(p.name for p in cfg.paths.inbox.iterdir()))


class MailboxUnrestricted(unittest.TestCase):
    def test_any_sender_accepted_when_restriction_off(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            cfg = make_cfg(root, "\nrestrict_senders = false\n")
            self.assertFalse(cfg.mail_in.restrict_senders)
            data = xlsx_bytes(root)
            box = mailbox.mbox(str(root / "box"))
            add_mail(box, "Bad <evil@evil.ru>", "ПОНМ", "any.xlsx", data)
            add_mail(box, "Bad <evil@evil.ru>", "ПОНМ", "program.exe", b"MZ")
            box.flush()
            box.close()

            mail_in.scan_mailbox(cfg)

            self.assertEqual(sorted(p.name for p in cfg.paths.inbox.iterdir()), ["any.xlsx"])

    def test_empty_list_allowed_only_without_restriction(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            base = CONFIG.replace('allowed_senders = ["good@company.ru", "@trusted.ru"]', "allowed_senders = []")
            path = root / "config.toml"
            path.write_text(base.format(root=root.as_posix()), encoding="utf-8")
            with self.assertRaises(ConfigError):
                load_config(path)
            path.write_text(base.format(root=root.as_posix()) + "\nrestrict_senders = false\n", encoding="utf-8")
            self.assertFalse(load_config(path).mail_in.restrict_senders)


class SmtpCertificate(unittest.TestCase):
    MAIL = """
    [mail_out]
    enabled = true
    from = "a@company.ru"
    to = ["b@company.ru"]
    [smtp]
    server = "10.50.40.58"
"""

    def test_verification_on_by_default_and_can_be_turned_off(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            self.assertTrue(make_cfg(root, self.MAIL).mail_out.smtp.verify_certificate)
            cfg = make_cfg(root, self.MAIL + "    verify_certificate = false\n")
            self.assertFalse(cfg.mail_out.smtp.verify_certificate)
            self.assertTrue(any("сертификат" in w for w in cfg.warnings))

    def test_context_does_not_verify_when_off(self):
        import ssl
        from unittest import mock
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            for flag, expected in (("true", ssl.CERT_REQUIRED), ("false", ssl.CERT_NONE)):
                smtp = make_cfg(root, self.MAIL + f"    verify_certificate = {flag}\n").mail_out.smtp
                with mock.patch("smtplib.SMTP") as fake:
                    mail_out._connect(smtp)
                    ctx = fake.return_value.starttls.call_args.kwargs["context"]
                    self.assertEqual(ctx.verify_mode, expected)


class MailFormat(unittest.TestCase):
    def _send(self, root: Path, src_name: str):
        import email
        import email.generator
        import io
        from unittest import mock
        from svodka.config import MailOut, Smtp
        att = root / "Сводка 2-х час ГПН-Ямал — Сводка ПОНМ октябрь_20261013_131928.xlsx"
        att.write_bytes(b"data")
        cfg = MailOut(True, "a@x.ru", ("b@x.ru",), False, True, 10, Smtp("h", 587, "none", True, "", ""))
        sent = {}
        with mock.patch.object(mail_out, "_connect") as conn:
            conn.return_value.send_message.side_effect = lambda m, **kw: sent.setdefault("m", m)
            mail_out.send_report(cfg, src_name, [att])
        buf = io.BytesIO()
        email.generator.BytesGenerator(buf, policy=sent["m"].policy.clone(linesep="\r\n")).flatten(sent["m"])
        return att.name, buf.getvalue()

    def test_long_russian_attachment_name_is_not_split(self):
        import email
        from email import policy
        with tempfile.TemporaryDirectory() as d:
            name, raw = self._send(Path(d), "Сводка ПОНМ октябрь.xlsb")
        # RFC 2231 'filename*0*=' клиенты и шлюзы не склеивают: вложение приходит как «Часть 1.2»
        self.assertNotIn(b"filename*", raw)
        self.assertNotIn(b"name*", raw)
        part = next(email.message_from_bytes(raw, policy=policy.default).iter_attachments())
        self.assertEqual(part.get_filename(), name)
        self.assertEqual(part.get_param("name"), name)

    def test_newline_in_subject_is_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(ValueError):
                self._send(Path(d), "a" + chr(10) + "Bcc: x@y.ru")


class CaptureBlock(unittest.TestCase):
    class Src:
        def __init__(self, bc, bn):
            self.cells = {"BC101": bc, "BN101": bn}

        def __getitem__(self, addr):
            return type("C", (), {"value": self.cells[addr]})()

    def test_ukpg_from_bc101_cps_from_bn101_by_day(self):
        from datetime import date
        from svodka import builder
        run = RunState()
        ws = Workbook().active
        ws["D111"], ws["D112"] = 2035, 0            # старые числа образца
        captured = {}
        for day, (bc, bn) in {1: (100.5, 1), 2: ("1 200,5", 2), 3: (300, "")}.items():
            d = date(2026, 10, day)
            captured[d] = builder.read_capture(self.Src(bc, bn), run)
        captured[date(2026, 9, 30)] = (999, 999)    # прошлый месяц в таблицу не попадает
        builder.fill_capture(ws, date(2026, 10, 3), captured, run)

        self.assertEqual((ws["D105"].value, ws["D107"].value), (300, None))        # текущие сутки (3-е число)
        self.assertEqual([ws.cell(111, c).value for c in (4, 5, 6, 7)], [100.5, 1200.5, 300, None])
        self.assertEqual([ws.cell(112, c).value for c in (4, 5, 6, 7)], [1, 2, None, None])

        builder.fill_capture(ws, date(2026, 10, 2), captured, run)                  # на 2-м числе 3-е не видно как «своё»
        self.assertEqual((ws["D105"].value, ws["D107"].value), (1200.5, 2))


class PdfPreparation(unittest.TestCase):
    def test_only_last_day_sheet_range_fit_to_one_page(self):
        from openpyxl import load_workbook
        from svodka import pdf
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            wb = Workbook()
            wb.active.title = "data_sms"
            for name in ("01.10", "02.10", "03.10"):
                wb.create_sheet(name)
            wb["03.10"]["B1"] = "сегодня"
            wb["03.10"].page_setup.scale = 46
            wb.create_sheet("смс")                        # служебный лист после дневных — не «последний»
            wb["смс"].sheet_state = "hidden"
            src, out = root / "src.xlsx", root / "out.xlsx"
            wb.save(src)

            title = pdf.prepare_single_page(src, out)

            self.assertEqual(title, "03.10")
            res = load_workbook(out)
            self.assertEqual([w.title for w in res.worksheets if w.sheet_state == "visible"], ["03.10"])
            ws = res["03.10"]
            self.assertEqual(ws.print_area, "'03.10'!$B$1:$P$108")
            self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
            self.assertEqual((ws.page_setup.fitToWidth, ws.page_setup.fitToHeight), (1, 1))
            self.assertIsNone(ws.page_setup.scale)
            self.assertEqual(ws["B1"].value, "сегодня")
            self.assertEqual(load_workbook(src)["01.10"].sheet_state, "visible")   # исходник не тронут


class Watcher(unittest.TestCase):
    def test_non_excel_file_goes_to_errors(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            cfg = make_cfg(root)
            cfg.paths.inbox.mkdir(parents=True)
            (cfg.paths.inbox / "fake.xlsx").write_text("не excel")
            watcher.scan_and_process(cfg)
            self.assertFalse((cfg.paths.inbox / "fake.xlsx").exists())
            self.assertTrue((cfg.paths.errors / "fake.xlsx").exists())

    def test_own_results_not_taken_as_source(self):
        self.assertFalse(watcher.is_source_candidate(Path("Сводка 2-х час ГПН-Ямал — a.xlsx")))
        self.assertFalse(watcher.is_source_candidate(Path("~$a.xlsx")))
        self.assertTrue(watcher.is_source_candidate(Path("Сводка ПОНМ октябрь.xlsb")))


class ManualValues(unittest.TestCase):
    def load(self, text):
        from svodka.manual import load_manual_values
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "manual_values.toml"
            p.write_text(textwrap.dedent(text), encoding="utf-8")
            return load_manual_values(p)

    def test_example_file_is_valid(self):
        from svodka.manual import load_manual_values
        cells = load_manual_values(Path("config/manual_values.example.toml"))
        self.assertEqual(cells["Q30"], 1004021.656)
        self.assertEqual((cells["Q55"], cells["R55"].strftime("%d.%m.%Y")), (0.67, "01.06.2026"))

    def test_missing_file_is_none(self):
        from svodka.manual import load_manual_values
        self.assertIsNone(load_manual_values(Path("нет-такого.toml")))

    def test_bp_and_coefficient_go_to_their_cells(self):
        cells = self.load("""
            [bp]
            "Конденсат ПНГ - БП" = 100
            [coefficients]
            "Стабильный конденсат ПГ (тёплый)" = { "Коэффициент" = 0.9, "Дата" = 2026-10-01 }
        """)
        self.assertEqual(cells["V52"], 100)
        self.assertEqual((cells["Q81"], cells["R81"].day, cells["R81"].month), (0.9, 1, 10))

    def test_mistakes_are_reported(self):
        for text in ('[bp]\n"Добыча ПНГ БП" = 1',                       # опечатка в названии
                     '[bp]\n"Добыча ПНГ - БП" = "1,5"',                   # число текстом
                     '[coefficients]\n"Стабильный конденсат ПНГ (УРГ)" = { "Коэффициент" = 0.4 }',
                     '[coefficients]\n"Стабильный конденсат ПНГ (УРГ)" = { "Коэффициент" = 0.4, "Дата" = "32.01.2026" }'):
            with self.subTest(text=text), self.assertRaises(ConfigError):
                self.load(text)


if __name__ == "__main__":
    unittest.main()
