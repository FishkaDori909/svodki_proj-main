"""Превращение готовой сводки в PDF через LibreOffice (если он установлен).

В PDF попадает не вся книга, а только последний лист сводки (текущие сутки), диапазон
PDF_AREA, уменьшенный до одной страницы. Сама сводка при этом не меняется: настройки
печати задаются на временной копии.
"""
import glob
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.pagebreak import ColBreak, RowBreak
from openpyxl.worksheet.properties import PageSetupProperties

PDF_AREA = "B1:P108"
_DAY_SHEET = re.compile(r"\d{2}\.\d{2}")
_CHAR_PT = 5.6   # примерная ширина одного знака ширины столбца Excel, в пунктах


class PdfConversionError(Exception):
    pass


_EXTRA_SOFFICE_GLOBS = (
    "/opt/libreoffice*/program/soffice",
    "/usr/lib/libreoffice*/program/soffice",
    "C:/Program Files/LibreOffice/program/soffice.exe",
)


def find_soffice():
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    for pattern in _EXTRA_SOFFICE_GLOBS:
        matches = sorted(glob.glob(pattern))
        if matches:
            return matches[-1]
    return None


def _last_sheet(wb):
    """Последний лист сводки: последний видимый лист с именем 'дд.мм', иначе последний видимый."""
    visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
    if not visible:
        raise PdfConversionError("в сводке нет видимых листов")
    days = [ws for ws in visible if _DAY_SHEET.fullmatch(ws.title)]
    return (days or visible)[-1]


def _area_is_wide(ws, area: str) -> bool:
    """True, если выбранный диапазон шире, чем выше (тогда страница альбомная)."""
    min_col, min_row, max_col, max_row = range_boundaries(area)
    width = sum((ws.column_dimensions[get_column_letter(c)].width or 8.43) * _CHAR_PT
                for c in range(min_col, max_col + 1)
                if not ws.column_dimensions[get_column_letter(c)].hidden)
    default_h = ws.sheet_format.defaultRowHeight or 15
    height = sum((ws.row_dimensions[r].height or default_h)
                 for r in range(min_row, max_row + 1) if not ws.row_dimensions[r].hidden)
    return width > height


def prepare_single_page(src: Path, dest: Path, area: str = PDF_AREA) -> str:
    """Сохраняет в dest копию книги, из которой в PDF уйдёт один лист: последний, диапазон area,
    вписанный в одну страницу. Возвращает название этого листа."""
    wb = load_workbook(src)
    target = _last_sheet(wb)
    for ws in wb.worksheets:
        # скрытые листы в PDF не попадают, а формулы, ссылающиеся на них, продолжают считаться
        ws.sheet_state = "visible" if ws is target else "hidden"
        ws.sheet_view.tabSelected = ws is target
    wb.active = wb.worksheets.index(target)

    target.print_area = area
    target.print_title_rows = None
    target.print_title_cols = None
    target.row_breaks = RowBreak()
    target.col_breaks = ColBreak()
    props = target.sheet_properties.pageSetUpPr or PageSetupProperties()
    props.fitToPage = True
    target.sheet_properties.pageSetUpPr = props
    target.page_setup.scale = None
    target.page_setup.fitToWidth = 1
    target.page_setup.fitToHeight = 1
    target.page_setup.orientation = "landscape" if _area_is_wide(target, area) else "portrait"
    target.print_options.horizontalCentered = True
    wb.save(dest)
    return target.title


def to_pdf(xlsx_path, out_dir=None, timeout: int = 120, area: str = PDF_AREA) -> Path:
    xlsx_path = Path(xlsx_path)
    out_dir = Path(out_dir) if out_dir else xlsx_path.parent
    out_dir.mkdir(parents=True, exist_ok=True)

    soffice = find_soffice()
    if not soffice:
        raise PdfConversionError("LibreOffice не установлен")

    with tempfile.TemporaryDirectory(prefix="lo_svodka_") as tmp:
        tmp = Path(tmp)
        prepared = tmp / xlsx_path.name        # то же имя -> PDF получит имя сводки
        try:
            prepare_single_page(xlsx_path, prepared, area)
        except PdfConversionError:
            raise
        except Exception as e:
            raise PdfConversionError(f"не удалось подготовить лист для PDF: {e}") from e

        cmd = [
            soffice, "--headless", "--norestore", "--invisible",
            f"-env:UserInstallation={(tmp / 'profile').as_uri()}",
            "--convert-to", "pdf", "--outdir", str(out_dir), str(prepared),
        ]
        try:
            result = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            raise PdfConversionError(f"LibreOffice не ответил за {timeout} с")
        except OSError as e:
            raise PdfConversionError(f"Не удалось запустить LibreOffice: {e}")

    pdf_path = out_dir / (xlsx_path.stem + ".pdf")
    if result.returncode != 0 or not pdf_path.exists():
        raise PdfConversionError(
            f"LibreOffice завершился с кодом {result.returncode}: "
            f"{(result.stderr or result.stdout).strip()[:300]}")
    return pdf_path
